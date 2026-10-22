import io
from datetime import date, timedelta

from openpyxl import load_workbook

from tests.factories import BANK_CSV, GROCERY_RECEIPT, receipt_pdf


def test_budget_status_and_threshold_notifications(user):
    cats = {c["name"]: c["id"] for c in user.get("/categories").json()}
    today = date.today()
    user.post("/budgets", {"name": "Food", "period": "monthly", "category_id": cats["Food and Dining"], "amount": "1000", "alert_threshold_pct": 80})
    user.post("/budgets", {"name": "Everything", "period": "monthly", "amount": "50000"})
    user.payment("850.00", "Swiggy", on=today.isoformat())
    status = {b["name"]: b for b in user.get("/budgets/status").json()}
    assert status["Food"]["spent"] == "850.00" and status["Food"]["state"] == "warning" and status["Food"]["remaining"] == "150.00"
    assert status["Everything"]["state"] == "ok"
    assert user.post("/notifications/run-checks").json()["created"] >= 1
    assert user.post("/notifications/run-checks").json()["created"] == 0  # deduplicated
    notes = user.get("/notifications").json()["items"]
    assert any(n["kind"] == "budget_threshold" for n in notes)
    user.post("/notifications/read-all")
    assert user.get("/notifications?unread=true").json()["total"] == 0


def test_recurring_predictions_do_not_count_until_paid(user):
    today = date.today()
    r = user.post("/recurring", {"name": "Netflix", "amount": "649", "frequency": "monthly", "next_due_date": (today + timedelta(days=2)).isoformat(), "is_subscription": True}).json()
    month = today.strftime("%Y-%m")
    assert user.dashboard(month)["gross_spending"] == "0.00"
    dash = user.get(f"/dashboard/summary?month={month}").json()
    assert dash["upcoming_recurring"][0]["predicted"] is True
    assert any(n["kind"] == "upcoming_bill" for n in user.get("/notifications").json()["items"]) or user.post("/notifications/run-checks").json()["created"] >= 1
    t = user.payment("649.00", "Netflix", on=(today + timedelta(days=2)).isoformat())
    paid = user.post(f"/recurring/{r['id']}/mark-paid", {"transaction_id": t["id"]}).json()
    assert paid["next_due_date"] > (today + timedelta(days=2)).isoformat()


def test_recurring_detection(user):
    for m in ("2026-06-05", "2026-07-05", "2026-08-06", "2026-09-05"):
        user.payment("499.00", "Spotify", on=m)
    found = user.post("/recurring/detect").json()["suggested"]
    assert found and found[0]["frequency"] == "monthly" and found[0]["status"] == "suggested"


def test_savings_goal(user):
    g = user.post("/savings-goals", {"name": "Laptop", "target_amount": "90000", "saved_amount": "30000", "target_date": (date.today() + timedelta(days=200)).isoformat()}).json()
    assert g["remaining"] == "60000.00" and g["monthly_needed"] is not None


def test_reports_match_dashboard_and_export_formats(user):
    bank = user.post("/accounts", {"name": "Bank", "type": "bank"}).json()
    imp = user.post("/imports", files={"file": ("s.csv", BANK_CSV.encode(), "text/csv")}, data={"account_id": bank["id"]}).json()
    user.post(f"/imports/{imp['id']}/confirm")
    user.bill("450.00", "ABC Restaurant")
    m = [c for c in user.candidates() if c["kind"] == "bill_payment"][0]
    user.post(f"/reconciliation/{m['id']}/confirm", {})
    rep = user.get("/reports/monthly?date_from=2026-10-01&date_to=2026-10-31").json()
    assert rep["meta"]["currency"] == "INR" and rep["meta"]["date_from"] == "2026-10-01"
    assert rep["totals"][3] == user.dashboard()["net_spending"] == "1345.00"

    csv = user.get("/reports/transactions?date_from=2026-10-01&date_to=2026-10-31&format=csv")
    assert csv.status_code == 200 and "attachment" in csv.headers["content-disposition"]
    text = csv.content.decode("utf-8-sig")
    assert "Period,2026-10-01 to 2026-10-31" in text and "Currency,INR" in text
    x = user.get("/reports/category?date_from=2026-10-01&date_to=2026-10-31&format=xlsx")
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["A1"].value == "Spending by category"
    pdf = user.get("/reports/reconciliation?date_from=2026-10-01&date_to=2026-10-31&format=pdf")
    assert pdf.content.startswith(b"%PDF")
    for kind in ("annual", "merchant", "person", "budget", "items", "loans", "unmatched", "duplicates"):
        assert user.get(f"/reports/{kind}?date_from=2026-01-01&date_to=2026-10-31").status_code == 200, kind


def test_csv_export_neutralises_formula_injection(user):
    user.payment("10.00", "Shop", description="=HYPERLINK(\"http://evil\")")
    text = user.get("/reports/transactions?date_from=2026-10-01&date_to=2026-10-31&format=csv").content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text


def test_end_to_end_receipt_statement_and_split(user):
    """Upload a bill, import the bank statement, reconcile, split a dinner, check every total."""
    bank = user.post("/accounts", {"name": "XYZ Bank", "type": "bank", "opening_balance": "25000.00", "opening_balance_date": "2026-10-01"}).json()
    doc = user.post("/documents", files={"file": ("freshmart.pdf", receipt_pdf(GROCERY_RECEIPT), "application/pdf")}).json()
    d = user.get(f"/documents/{doc['id']}").json()["extraction"]["fields"]
    user.post(
        f"/documents/{doc['id']}/confirm",
        {"merchant_name": d["merchant_name"]["value"], "occurred_on": d["date"]["value"], "amount": d["total"]["value"], "payment_reference": d["payment_reference"]["value"], "doc_kind": "invoice"},
    )
    imp = user.post("/imports", files={"file": ("s.csv", BANK_CSV.encode(), "text/csv")}, data={"account_id": bank["id"]}).json()
    user.post(f"/imports/{imp['id']}/confirm")
    high = [c for c in user.candidates() if c["confidence"] == "high"]
    assert len(high) == 1  # the grocery bill <-> FRESHMART UPI row (same reference)
    r = user.post("/reconciliation/bulk-confirm", {"match_ids": [high[0]["id"]], "confirm": True})
    assert r.json()["confirmed"] == 1

    rahul = user.post("/counterparties", {"name": "Rahul"}).json()
    user.post(
        "/shared-expenses",
        {"description": "Pizza night", "total_amount": "900", "occurred_on": "2026-10-22", "payer_counterparty_id": rahul["id"], "participants": [{"counterparty_id": None}, {"counterparty_id": rahul["id"]}]},
    )
    t = user.dashboard()
    # 450 restaurant UPI + 595 groceries (bill, payment linked) + 300 to Rahul + 450 my pizza share
    assert t["gross_spending"] == "1795.00"
    assert t["income"] == "85000.00"
    bal = user.get(f"/accounts/{bank['id']}/balance").json()["balance"]
    assert bal == "96655.00"  # matches the statement's closing balance: no double debit
    people = {p["name"]: p for p in user.get("/counterparties/balances").json()}
    assert people["Rahul"]["payable"] == "450.00"
