"""Bank / card / UPI / wallet statement parsing: table extraction, column detection, row normalisation."""
import csv
import hashlib
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from app.modules.merchants.normalize import parse_narration

HEADER_HINTS = {
    "date": [r"^(txn|transaction|tran|posting|value)?\s*date$", r"^date$", r"txn\s*date", r"transaction\s*date", r"^dt$"],
    "description": [r"narration", r"description", r"particulars", r"details", r"remarks", r"transaction\s*details", r"merchant", r"^payee", r"^paid\s*to"],
    "debit": [r"withdrawal", r"^debit", r"\bdr\b", r"debit\s*amount", r"amount\s*\(?dr\)?", r"paid\s*out", r"money\s*out"],
    "credit": [r"deposit", r"^credit", r"\bcr\b", r"credit\s*amount", r"amount\s*\(?cr\)?", r"paid\s*in", r"money\s*in"],
    "amount": [r"^amount", r"transaction\s*amount", r"^amt", r"amount\s*\((inr|rs)"],
    "direction": [r"^(dr|cr)\s*/\s*(cr|dr)$", r"^type$", r"debit\s*/\s*credit", r"^dr/cr$", r"transaction\s*type", r"^status$"],
    "reference": [r"ref", r"utr", r"chq", r"cheque", r"transaction\s*id", r"txn\s*id", r"upi\s*ref", r"rrn"],
    "balance": [r"balance", r"closing"],
    "currency": [r"^currency$", r"^ccy$"],
}
DATE_FORMATS = ["%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y", "%Y-%m-%d", "%d %b %Y", "%d-%b-%Y", "%d %b %y", "%d-%b-%y", "%b %d, %Y", "%m/%d/%Y", "%d %B %Y", "%Y/%m/%d"]


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "")).strip().lower()


def _header_score(row: list) -> int:
    score = 0
    for cell in row:
        c = _norm(cell)
        if not c:
            continue
        if any(re.search(p, c) for pats in HEADER_HINTS.values() for p in pats):
            score += 1
    return score


def _find_header(rows: list[list]) -> int:
    best, best_i = 0, 0
    for i, r in enumerate(rows[:30]):
        sc = _header_score(r)
        if sc > best:
            best, best_i = sc, i
    if best < 2:
        raise ValueError("Could not find a header row with date/description/amount columns")
    return best_i


def read_table(fmt: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    if fmt == "csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect)]
    elif fmt == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else (v.strftime("%Y-%m-%d") if isinstance(v, (datetime, date)) else str(v)) for v in r])
            if len(rows) > 20000:
                break
    elif fmt == "pdf":
        import pdfplumber

        rows = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:50]:
                for table in page.extract_tables():
                    rows.extend([[c or "" for c in r] for r in table])
        if not rows:
            raise ValueError("No tables were found in this PDF statement. Export a CSV or XLSX from your bank instead.")
    else:
        raise ValueError("Unsupported format")
    rows = [[str(c).strip() for c in r] for r in rows if any(str(c).strip() for c in r)]
    h = _find_header(rows)
    headers = [str(c).strip() or f"Column {i + 1}" for i, c in enumerate(rows[h])]
    body = []
    for r in rows[h + 1 :]:
        if _header_score(r) >= 3 and [_norm(c) for c in r] == [_norm(c) for c in rows[h]]:
            continue  # repeated header on a later PDF page
        body.append((r + [""] * len(headers))[: len(headers)])
    return headers, body


def detect_mapping(headers: list[str]) -> dict:
    mapping: dict[str, str | None] = {k: None for k in HEADER_HINTS}
    used = set()
    for field in ["date", "description", "direction", "debit", "credit", "amount", "reference", "balance", "currency"]:
        for h in headers:
            if h in used:
                continue
            c = _norm(h)
            if any(re.search(p, c) for p in HEADER_HINTS[field]):
                if field == "date" and "value" in c and any("value" not in _norm(x) and re.search(r"date", _norm(x)) for x in headers if x != h):
                    continue  # prefer transaction date over value date
                mapping[field] = h
                used.add(h)
                break
    if mapping["debit"] and mapping["credit"]:
        mapping["amount"] = None
    return mapping


def parse_amount(s: str) -> Decimal | None:
    if s is None:
        return None
    t = str(s).strip()
    if not t or t in ("-", "--"):
        return None
    neg = t.startswith("(") and t.endswith(")")
    t = re.sub(r"[₹,\s()]|rs\.?|inr", "", t, flags=re.I)
    suffix = None
    m = re.match(r"^(-?\d+(?:\.\d+)?)\s*(dr|cr)?$", t, re.I)
    if not m:
        return None
    try:
        v = Decimal(m.group(1))
    except InvalidOperation:
        return None
    suffix = (m.group(2) or "").lower()
    if neg or suffix == "dr":
        v = -abs(v)
    return v


def _try(v: str, fmt: str) -> date | None:
    for cand in (v, v.split(" ")[0], v[:10]):
        try:
            return datetime.strptime(cand, fmt).date()
        except (ValueError, TypeError):
            continue
    return None


def infer_date_format(values: list[str]) -> str | None:
    """Picks the first format that parses every sample (so 03/04/2026 is read consistently for the whole file)."""
    samples = [v.strip() for v in values if v and v.strip()][:200]
    if not samples:
        return None
    for fmt in DATE_FORMATS:
        if all(_try(v, fmt) for v in samples):
            return fmt
    return None


def parse_date(v: str, fmt: str | None) -> date | None:
    v = (v or "").strip()
    if not v:
        return None
    for f in [fmt] if fmt else DATE_FORMATS:
        d = _try(v, f)
        if d:
            return d
    return None


CC_PAYMENT_RE = re.compile(r"credit\s*card\s*(bill|payment)|cc\s*(bill|payment|pymt)|card\s*payment|payment\s*received|thank\s*you\s*for\s*(your\s*)?payment", re.I)
SELF_TRANSFER_RE = re.compile(r"self\s*transfer|own\s*account|to\s*self|sweep|fd\s*booking|transfer\s*to\s*savings", re.I)
REFUND_RE = re.compile(r"refund|reversal|reversed|\brev\b|chargeback|cashback", re.I)


def normalise_rows(raw_rows: list[list[str]], headers: list[str], mapping: dict, date_format: str | None, currency: str, source_kind: str) -> list[dict]:
    idx = {h: i for i, h in enumerate(headers)}

    def col(row, key):
        h = mapping.get(key)
        return row[idx[h]] if h in idx else ""

    fmt = date_format or infer_date_format([col(r, "date") for r in raw_rows])
    out = []
    seen: dict[str, int] = {}
    for n, row in enumerate(raw_rows):
        errors = []
        d = parse_date(col(row, "date"), fmt)
        if d is None:
            if not any(row) or not col(row, "date"):
                continue  # blank / summary line without a date
            errors.append(f"Unrecognised date '{col(row, 'date')}'")
        desc = re.sub(r"\s+", " ", col(row, "description")).strip()
        amount = None
        direction = None
        if mapping.get("debit") or mapping.get("credit"):
            dr = parse_amount(col(row, "debit")) if mapping.get("debit") else None
            cr = parse_amount(col(row, "credit")) if mapping.get("credit") else None
            if dr and dr != 0 and (not cr or cr == 0):
                amount, direction = abs(dr), "out"
            elif cr and cr != 0 and (not dr or dr == 0):
                amount, direction = abs(cr), "in"
            elif dr and cr:
                errors.append("Both debit and credit have values")
        if amount is None and mapping.get("amount"):
            a = parse_amount(col(row, "amount"))
            if a is not None and a != 0:
                amount = abs(a)
                flag = _norm(col(row, "direction")) if mapping.get("direction") else ""
                if flag:
                    if re.search(r"\b(dr|debit|withdrawal|paid|sent|debited)\b", flag):
                        direction = "out"
                    elif re.search(r"\b(cr|credit|deposit|received|credited)\b", flag):
                        direction = "in"
                    elif re.search(r"failed|declined", flag):
                        errors.append("Transaction marked failed")
                if direction is None:
                    if source_kind in ("credit_card", "upi"):
                        direction = "out" if a > 0 else "in"  # card/UPI exports list spends as positive numbers
                    else:
                        direction = "out" if a < 0 else "in"  # bank/wallet: signed amounts
        if amount is None and not errors:
            errors.append("No amount found")
        cur = (col(row, "currency") or currency).strip().upper()[:3] or currency
        parsed = parse_narration(desc)
        ref = (col(row, "reference") or "").strip() or parsed["reference"]
        ref = re.sub(r"\s+", "", ref)[:64] if ref else None
        if ref and set(ref) <= {"0"}:
            ref = None
        base_fp = f"{d}|{direction}|{amount}|{re.sub(r'[^a-z0-9]', '', desc.lower())}|{ref or ''}"
        occurrence = seen.get(base_fp, 0)
        seen[base_fp] = occurrence + 1
        tx_type = "expense" if direction == "out" else "income"
        if direction == "in" and REFUND_RE.search(desc):
            tx_type = "refund"
        if CC_PAYMENT_RE.search(desc) and (source_kind == "credit_card" and direction == "in" or source_kind == "bank" and direction == "out"):
            tx_type = "credit_card_payment"
        elif SELF_TRANSFER_RE.search(desc):
            tx_type = "transfer"
        out.append(
            {
                "row_index": n,
                "raw": dict(zip(headers, row)),
                "occurred_on": d.isoformat() if d else None,
                "description": desc[:500],
                "amount": format(amount, "f") if amount is not None else None,
                "direction": direction,
                "currency": cur,
                "reference": ref,
                "balance": format(parse_amount(col(row, "balance")), "f") if mapping.get("balance") and parse_amount(col(row, "balance")) is not None else None,
                "payee": parsed["payee"],
                "vpa": parsed["vpa"],
                "channel": parsed["channel"],
                "suggested_type": tx_type,
                "occurrence": occurrence,
                "base_fingerprint": base_fp,
                "errors": errors,
            }
        )
    return out


def fingerprint(household_id, account_id, row: dict) -> str:
    """Stable across re-imports of the same statement; distinct for genuinely repeated identical rows."""
    payload = f"{household_id}|{account_id or ''}|{row['base_fingerprint']}|{row['occurrence']}"
    return hashlib.sha256(payload.encode()).hexdigest()
